"""Regression examples reproduce structures observed in real inputs."""

import csv
import json
from datetime import date
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook
from yoyoflot.readers import (
    boarding_csv,
    forum_json,
    pointz_xml,
    skyteam_yaml,
    boardingpass_xlsx,
    sirena_fixed,
)


def collect(reader, path, **kwargs):
    out = {}
    for batch in reader(path, batch_size=2, **kwargs):
        batch.to_arrow()
        out.setdefault(batch.table_name, []).extend(batch.records)
    return out


def test_real_csv_us_birth_date(tmp_path):
    path = tmp_path / "boarding.csv"
    row = dict.fromkeys(boarding_csv.EXPECTED_COLUMNS, "")
    row.update(
        PassengerFirstName="A", PassengerBirthDate="03/10/1983", TicketNumber="000123"
    )
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(boarding_csv.EXPECTED_COLUMNS)
        w.writerow([row[k] for k in boarding_csv.EXPECTED_COLUMNS])
    assert (
        collect(boarding_csv.read, path)["person_observations"][0]["birth_date"]
        == "1983-03-10"
    )


def test_real_xml_name_uid_composite_card(tmp_path):
    path = tmp_path / "pointz.xml"
    path.write_text(
        '<PointzAggregatorUsers><user uid="001"><name first="IVAN" last="IVANOV"/><cards><card number="FB 0012"><bonusprogramm>Flying Blue</bonusprogramm><activities><activity type="Flight"><Code>AB1</Code><Date>2017-01-01</Date></activity></activities></card></cards></user></PointzAggregatorUsers>',
        encoding="utf-8",
    )
    d = collect(pointz_xml.read, path)
    assert d["person_observations"][0]["name_raw"] == "IVAN IVANOV"
    assert d["person_observations"][0]["external_id"] == "001"
    assert (d["identifiers"][0]["programme"], d["identifiers"][0]["value"]) == (
        "FB",
        "0012",
    )
    assert (
        d["travel_observations"][0]["loyalty_identifier_id"]
        == d["identifiers"][0]["identifier_id"]
    )


def test_real_json_null_passports_name_programm(tmp_path):
    path = tmp_path / "forum.json"
    path.write_text(
        json.dumps(
            {
                "Forum Profiles": [
                    {
                        "Real Name": {"First Name": "IVAN", "Last Name": "IVANOV"},
                        "Travel Documents": [{"Passports": None}],
                        "Loyality Programm": [
                            {"programm": "SU", "Number": " 0012", "Status": "Elite"}
                        ],
                        "Registered Flights": [
                            {"Date": "2017-01-01", "Flight": "AB1", "Codeshare": True}
                        ],
                    },
                    {"Real Name": {"First Name": None, "Last Name": None}},
                ]
            }
        ),
        encoding="utf-8",
    )
    d = collect(forum_json.read, path)
    assert (
        d["person_observations"][0]["name_raw"] == "IVAN IVANOV"
        and d["person_observations"][1]["name_raw"] is None
    )
    assert (d["identifiers"][0]["programme"], d["identifiers"][0]["value"]) == (
        "SU",
        "0012",
    )
    assert d["travel_observations"][0]["codeshare"] == "True"
    assert not d.get("errors")


def test_real_yaml_composite_keys(tmp_path):
    path = tmp_path / "sky.yaml"
    path.write_text(
        "'2017-01-01':\n  AB001:\n    FF:\n      FB 0012: {CLASS: Y, FARE: YRST}\n      SU 0012: {CLASS: J, FARE: JRST}\n    FROM: AAA\n    TO: BBB\n    STATUS: LANDED\n",
        encoding="utf-8",
    )
    d = collect(skyteam_yaml.read, path)
    assert {(r["programme"], r["value"]) for r in d["identifiers"]} == {
        ("FB", "0012"),
        ("SU", "0012"),
    }
    assert d["travel_observations"][0]["fare_raw"] == "YRST"
    assert d["travel_observations"][0]["status_raw"] == "LANDED"
    assert d["travel_observations"][0]["event_kind"] == "exchange_record"
    assert not d.get("errors")


def test_real_card_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    for coord, value in boardingpass_xlsx.CARD_LABELS.items():
        ws[coord] = value
    for coord, value in {
        "B3": "IVAN IVANOV",
        "F3": "FB 0012",
        "A5": "AB001",
        "D5": "ALPHA",
        "H5": "BETA",
        "D7": "AAA",
        "H7": "BBB",
        "A9": "2017-01-01",
        "C9": "23:30",
        "B13": "ABCDEF",
        "E13": "000123",
    }.items():
        ws[coord] = value
    book = tmp_path / "2018-01-01.xlsx"
    wb.save(book)
    wb.close()
    path = tmp_path / "boarding.zip"
    with ZipFile(path, "w") as z:
        z.write(book, book.name)
    cov = tmp_path / "cov.csv"
    d = collect(boardingpass_xlsx.read, path, coverage_path=cov)
    assert len(d["person_observations"]) == len(d["travel_observations"]) == 1
    assert d["travel_observations"][0]["departure_local"] == "2017-01-01T23:30"
    assert (
        d["travel_observations"][0]["origin_raw"] == "AAA"
        and d["travel_observations"][0]["ticket_number"] == "000123"
    )
    assert d["identifiers"][1]["programme"] == "FB"
    with cov.open(encoding="utf-8") as f:
        row = next(csv.DictReader(f))
    assert row["header_version"] == "boarding_card_v1" and row["accepted_rows"] == "1"


def test_excel_date_style_keeps_date_precision(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["PassengerName", "FlightDate", "TicketNumber"])
    ws.append(["X", date(2017, 1, 1), "000123"])
    book = tmp_path / "date.xlsx"
    wb.save(book)
    wb.close()
    path = tmp_path / "boarding.zip"
    with ZipFile(path, "w") as z:
        z.write(book, book.name)
    d = collect(boardingpass_xlsx.read, path)
    assert (
        d["travel_observations"][0]["departure_local"] == "2017-01-01"
        and d["travel_observations"][0]["departure_precision"] == "date"
    )


def test_confirmed_sirena_ranges(tmp_path):
    import yaml

    config = Path(__file__).parents[1] / "configs/sirena_columns.yaml"
    cfg = yaml.safe_load(config.read_text(encoding="utf-8"))
    values = {
        "PaxName": "ИВАН ИВАНОВ",
        "PaxBirthDate": "N/A",
        "DepartDate": "2017-01-01",
        "DepartTime": "23:30",
        "ArrivalDate": "2017-01-02",
        "ArrivalTime": "01:00",
        "Flight": "AB001",
        "CodeSh": "NO",
        "From": "AAA",
        "Dest": "BBB",
        "Code": "ABCDEF",
        "e-Ticket": "000123",
        "TravelDoc": "0012 345678",
        "PaxAdditionalInfo": "FF#FB 0012",
    }
    line = "".join(
        values.get(c["name"], "").ljust(c["end"] - c["start"]) for c in cfg["columns"]
    )
    path = tmp_path / "sirena.tab"
    path.write_text(cfg["expected_header"] + "\n" + line + "\n", encoding="utf-8")
    d = collect(sirena_fixed.read, path, columns=config)
    assert d["person_observations"][0]["birth_date"] is None
    assert (
        d["travel_observations"][0]["booking_code"] == "ABCDEF"
        and d["travel_observations"][0]["codeshare"] == "NO"
    )
    assert d["travel_observations"][0]["arrival_local"] == "2017-01-02T01:00"
    assert d["identifiers"][2]["value"] == "0012"
    assert not d.get("errors")


def test_nested_passport_locators_keep_container_and_index(tmp_path):
    path = tmp_path / "forum.json"
    path.write_text(
        json.dumps(
            {
                "Forum Profiles": [
                    {
                        "Real Name": "X",
                        "Travel Documents": [
                            {"Passports": None},
                            {"Passports": [{"Number": "0001"}, {"Number": "0002"}]},
                        ],
                        "Registered Flights": [
                            {"Date": "2017-01-01"} for _ in range(3)
                        ],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    d = collect(forum_json.read, path)
    assert len(d["identifiers"]) == 2 and len(d["travel_observations"]) == 3
    assert d["identifiers"][0]["locator"].endswith("Travel Documents/1/Passports/0")
    assert d["identifiers"][1]["locator"].endswith("Travel Documents/1/Passports/1")
