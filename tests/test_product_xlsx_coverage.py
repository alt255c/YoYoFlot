"""Coverage admission uses actual reader layouts and canonical locator evidence."""
import csv
import json
import shutil
from pathlib import Path
from zipfile import ZipFile

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from yoyoflot.io import ingest
from yoyoflot.product.coverage import FIELDS, validate_xlsx_coverage
from yoyoflot.product.input import verify
from yoyoflot.provenance import fingerprint
from yoyoflot.readers.boardingpass_xlsx import CARD_LABELS, read
from tests.test_pipeline import args
from tests.test_product_provenance import refresh_t16


def csv_rows(path):
    with path.open(encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))


def write_csv(path, rows):
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)


@pytest.fixture
def evidence(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Пассажиры'
    sheet.append(['PassengerName', 'FlightNumber', 'FlightDate', 'TicketNumber'])
    sheet.append(['Контроль А', 'SU1', '2017-01-01', 123])  # accepted + warning
    sheet.append(['Контроль Б', '=1+1', '2017-01-02', 'T2'])  # accepted + cell error
    sheet.append(['Контроль В', 'SU3', '2017-01-03', 'T3', 'excess'])  # rejected row
    book.create_sheet('Empty')
    unsupported = book.create_sheet('Lookup')
    unsupported.append(['lookup', 'value'])
    unsupported.append(['a', 'b'])
    unsupported.append(['c', 'd'])
    bad_card = book.create_sheet('BadCard')
    bad_card['A1'] = 'BOARDING PASS'
    card = book.create_sheet('Card')
    for coord, value in CARD_LABELS.items():
        card[coord] = value
    for coord, value in {'B3': 'Контроль Г', 'A5': 'SU4', 'A9': '2017-01-04',
                         'D7': 'SVO', 'H7': 'LED', 'B13': 'PNR4', 'E13': 'T4'}.items():
        card[coord] = value
    workbook = tmp_path / 'book.xlsx'
    book.save(workbook)
    raw = tmp_path / 'boarding.zip'
    with ZipFile(raw, 'w') as archive:
        archive.write(workbook, 'nested/Книга.xlsx')
        archive.writestr('broken.xlsx', b'not OOXML')
        archive.writestr('README.txt', b'archive note')
    output = tmp_path / 'reader'
    report = ingest(read, raw, 'boardingpass_xlsx', output, batch_size=2)
    # Keep the actual reader's locators and counts; the validator reads only
    # those canonical projections, independent of normalization formulas.
    tables = tmp_path / 'tables'
    tables.mkdir()
    balances = []
    for source, target in [('person_observations', 'observations'),
                           ('travel_observations', 'travel_observations'), ('errors', 'errors')]:
        shutil.copyfile(output / (source + '.parquet'), tables / (target + '.parquet'))
        n = report['counts'][source]
        if n:
            balances.append(dict(source_id='boardingpass_xlsx', entity_table=target,
                                 input_rows=n, output_rows=n, duplicates_removed=0))
    source_path = str(output / 'source_report.json')
    canonical = {'source_reports': {source_path: report},
                 'inputs': {source_path: fingerprint(Path(source_path))}, 'balances': balances}
    return dict(path=output / 'workbook_coverage.csv', report=report, tables=tables,
                canonical=canonical, raw={'path': str(raw), 'sha256': fingerprint(raw), 'bytes': raw.stat().st_size})


def validate(evidence):
    return validate_xlsx_coverage(evidence['path'], evidence['report'],
                                 canonical_report=evidence['canonical'], canonical_tables=evidence['tables'],
                                 raw_input=evidence['raw'])


def test_actual_layouts_partial_errors_and_unicode_are_valid(evidence):
    result = validate(evidence)
    assert result['status'] == 'passed'
    assert result['coverage_rows'] == 7 and result['workbooks'] == 2 and result['sheets'] == 5
    assert result['input_rows'] == 7 and result['accepted_rows'] == 3
    assert result['errors'] == 6 and result['excluded_rows'] == 2
    assert result['boarding_card_sheets'] == 2
    assert result['canonical_evidence'] == {
        'observations': 3, 'travel_observations': 3, 'errors': 6, 'per_sheet_counts_verified': True}
    assert evidence['report']['status'] == 'complete_with_errors'
    assert result['coverage_sha256'] == fingerprint(evidence['path'])
    assert 'Книга.xlsx' in evidence['path'].read_text()


@pytest.mark.parametrize('mutation', ['drop_sheet', 'duplicate_sheet', 'replace_sheet', 'accepted',
                                    'negative', 'unknown_layout', 'unknown_status', 'missing_cell', 'extra_cell'])
def test_malformed_or_incomplete_coverage_is_rejected(evidence, mutation):
    path = evidence['path']
    rows = csv_rows(path)
    if mutation == 'drop_sheet': rows.pop(0)
    elif mutation == 'duplicate_sheet': rows.append(dict(rows[0]))
    elif mutation == 'replace_sheet': rows[0]['sheet'] = 'AbsentInCanonical'
    elif mutation == 'accepted': rows[0]['accepted_rows'] = '1'
    elif mutation == 'negative': rows[0]['input_rows'] = '-1'
    elif mutation == 'unknown_layout': rows[0]['header_version'] = 'invented'
    elif mutation == 'unknown_status': rows[0]['status'] = 'complete'
    write_csv(path, rows)
    if mutation == 'missing_cell':
        lines = path.read_text().splitlines()
        lines[-1] = ','.join(lines[-1].split(',')[:-1])
        path.write_text('\n'.join(lines) + '\n')
    elif mutation == 'extra_cell':
        with path.open('a') as stream: stream.write('a,b,c,1,1,0,ok,,extra\n')
    with pytest.raises(ValueError, match='XLSX coverage'):
        validate(evidence)


@pytest.mark.parametrize('header', [','.join(FIELDS[:-1]), ','.join(FIELDS + ('reason',)), 'wrong,header'])
def test_schema_is_required_even_for_empty_csv(evidence, header):
    evidence['path'].write_text(header + '\n')
    with pytest.raises(ValueError, match='schema'):
        validate(evidence)


@pytest.mark.parametrize('key', ['workbooks', 'sheets', 'input_rows', 'boarding_card_sheets', 'excluded_rows'])
def test_all_reader_counters_are_bound(evidence, key):
    evidence['report']['input_counts'][key] += 1
    with pytest.raises(ValueError, match='total mismatch'):
        validate(evidence)


def test_coverage_totals_alone_do_not_allow_reassigned_canonical_evidence(evidence):
    path = evidence['tables'] / 'travel_observations.parquet'
    table = pq.read_table(path)
    rows = table.to_pylist()
    rows[0]['locator'] = rows[0]['locator'].replace('/Пассажиры/', '/Empty/')
    pq.write_table(pa.Table.from_pylist(rows, schema=table.schema), path)
    with pytest.raises(ValueError, match='differs for workbook/sheet'):
        validate(evidence)


def test_missing_canonical_evidence_rejected(evidence):
    (evidence['tables'] / 'observations.parquet').unlink()
    with pytest.raises(ValueError, match='missing canonical evidence'):
        validate(evidence)


def test_raw_inventory_and_embedded_report_are_bound(evidence):
    evidence['raw']['sha256'] = '0' * 64
    with pytest.raises(ValueError, match='raw inventory'):
        validate(evidence)
    evidence['raw']['sha256'] = evidence['report']['sha256']
    key = next(iter(evidence['canonical']['source_reports']))
    evidence['canonical']['source_reports'][key] = dict(evidence['report'], seconds=999)
    with pytest.raises(ValueError, match='canonical source report mismatch'):
        validate(evidence)


def test_canonical_duplicate_input_copies_are_accounted_not_rejected(evidence):
    for balance in evidence['canonical']['balances']:
        balance['duplicates_removed'] = balance['input_rows']
        balance['input_rows'] *= 2
    assert validate(evidence)['canonical_evidence']['observations'] == 3


def test_empty_archive_has_valid_empty_coverage(tmp_path):
    raw = tmp_path / 'empty.zip'
    with ZipFile(raw, 'w'): pass
    report = ingest(read, raw, 'boardingpass_xlsx', tmp_path / 'reader')
    result = validate_xlsx_coverage(tmp_path / 'reader/workbook_coverage.csv', report)
    assert result['coverage_rows'] == result['workbooks'] == result['accepted_rows'] == 0


def test_duplicate_key_is_detected_across_streaming_batches(tmp_path):
    path = tmp_path / 'coverage.csv'
    rows = [dict(workbook='one.xlsx', sheet=f'Sheet{i}', header_version='', input_rows=0,
                 accepted_rows=0, errors=0, status='empty', reason='') for i in range(8200)]
    write_csv(path, rows + [rows[0]])
    report = dict(source_id='boardingpass_xlsx', input_path='archive.zip', sha256='1' * 64,
                  contract_version='1.0', status='complete', error_severities={},
                  counts=dict(person_observations=0, travel_observations=0, errors=0),
                  input_counts=dict(workbooks=1, sheets=8201))
    with pytest.raises(ValueError, match='duplicate workbook/sheet'):
        validate_xlsx_coverage(path, report)


def test_checksum_valid_truncation_cannot_enter_product(args):
    from yoyoflot.pipeline import run

    run(**args)
    root = args['output'] / 'tables'
    verify(root)
    coverage = args['output'] / 'sources/boardingpass_xlsx/workbook_coverage.csv'
    original = coverage.read_bytes()
    rows = csv_rows(coverage)
    assert rows
    write_csv(coverage, rows[:-1])
    refresh_t16(args['output'])  # Valid checksum is deliberately insufficient.
    with pytest.raises(ValueError, match='XLSX coverage'):
        verify(root)
    coverage.write_bytes(original)
    refresh_t16(args['output'])
    verify(root)
    source_report = coverage.parent / 'source_report.json'
    original_report = source_report.read_bytes()
    changed = json.loads(original_report)
    changed['seconds'] += 1
    source_report.write_text(json.dumps(changed))
    refresh_t16(args['output'])
    with pytest.raises(ValueError, match='source report fingerprint differs from canonical'):
        verify(root)
    source_report.write_bytes(original_report)
    coverage.unlink()
    refresh_t16(args['output'])
    with pytest.raises(ValueError, match='Missing T16 XLSX coverage'):
        verify(root)
