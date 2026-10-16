import csv
import json
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from yoyoflot.io import ingest
from yoyoflot.readers import (
    boarding_csv,
    boardingpass_xlsx,
    forum_json,
    pointz_xml,
    sirena_fixed,
    skyteam_yaml,
)


def collect(reader, path, **kw):
    out = {}
    for batch in reader(path, batch_size=2, **kw):
        batch.to_arrow()
        assert len(batch.records) <= 2
        out.setdefault(batch.table_name, []).extend(batch.records)
    return out


def write_json(tmp_path, profiles):
    path = tmp_path / "forum.json"
    path.write_text(
        json.dumps({"Forum Profiles": profiles}, ensure_ascii=False), encoding="utf-8"
    )
    return path


def test_csv_errors_locators_quotes_and_dates(tmp_path):
    path = tmp_path / "boarding.csv"
    row = dict.fromkeys(boarding_csv.EXPECTED_COLUMNS, "")
    row.update(
        PassengerFirstName='Анна "А"',
        PassengerDocument="000123",
        TicketNumber="0001",
        FlightDate="2017-01-01",
        FlightTime="23:30",
        Destination="BBB",
    )
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(boarding_csv.EXPECTED_COLUMNS)
        w.writerow([row[k] for k in boarding_csv.EXPECTED_COLUMNS])
        w.writerow([])
        w.writerow(["wrong", "count"])
        row["FlightDate"] = "2017-02-30"
        w.writerow([row[k] for k in boarding_csv.EXPECTED_COLUMNS])
    data = collect(boarding_csv.read, path)
    assert len(data["person_observations"]) == len(data["travel_observations"]) == 2
    assert data["person_observations"][0]["locator"] == "line:2-2"
    assert data["person_observations"][1]["locator"] == "line:5-5"
    assert data["travel_observations"][0]["departure_local"] == "2017-01-01T23:30"
    assert data["travel_observations"][1]["departure_local"] is None
    assert data["travel_observations"][0]["origin_raw"] is None
    assert data["identifiers"][0]["value"] == "000123"
    assert {r["reason"] for r in data["errors"]} == {
        "blank_row",
        "column_count:2",
        "invalid_departure_date",
    }
    assert len({r["ticket_number"] for r in data["travel_observations"]}) == 1


def test_json_no_cartesian_multiplication_and_empty_profile(tmp_path):
    path = write_json(
        tmp_path,
        {
            "p/1": {
                "Real Name": "Иван",
                "NickName": "abc",
                "Travel Documents": {
                    "Passports": [{"Number": "001"}, {"Number": "002"}]
                },
                "Loyality Programm": {"AA": "0012", "BB": "0012"},
                "Registered Flights": [
                    {
                        "Flight": "AB001",
                        "Date": "2017-01-01",
                        "Departure": {"City": "A"},
                    }
                    for _ in range(3)
                ],
            },
            "p2": {"NickName": "n", "Registered Flights": []},
        },
    )
    data = collect(forum_json.read, path)
    assert len(data["person_observations"]) == 2
    assert len(data["identifiers"]) == 4
    assert len(data["travel_observations"]) == 3
    assert data["person_observations"][1]["name_raw"] is None
    assert {
        (r["programme"], r["value"])
        for r in data["identifiers"]
        if r["kind"] == "loyalty_card"
    } == {("AA", "0012"), ("BB", "0012")}
    assert "p~11" in data["person_observations"][0]["locator"]
    assert data["travel_observations"][0]["departure_precision"] == "date"
    assert data["travel_observations"][0]["departure_local"] == "2017-01-01"


def test_json_invalid_syntax_is_fatal(tmp_path):
    path = tmp_path / "bad.json"
    path.write_text('{"Forum Profiles":[{"Real Name":"X"},', encoding="utf-8")
    with pytest.raises(Exception):
        ingest(forum_json.read, path, "forum_json", tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_xml_namespace_cards_activities_and_no_documents(tmp_path):
    path = tmp_path / "pointz.xml"
    path.write_text(
        '<PointzAggregatorUsers xmlns="urn:test"><user id="001"><name>A &amp; B</name><cards><card number="0012"><bonusprogramm>AA</bonusprogramm><activities><activity><Code>AB1</Code><Date>2017-01-01</Date></activity></activities></card><card number="0012"><bonusprogramm>BB</bonusprogramm></card></cards><activities><activity><Arrival>BBB</Arrival></activity><activity><Arrival>BBB</Arrival></activity></activities></user><user><name>Empty</name><cards/></user></PointzAggregatorUsers>',
        encoding="utf-8",
    )
    data = collect(pointz_xml.read, path)
    assert len(data["person_observations"]) == 2
    assert len(data["identifiers"]) == 2
    assert len(data["travel_observations"]) == 3
    assert len({r["travel_id"] for r in data["travel_observations"]}) == 3
    assert data["person_observations"][0]["name_raw"] == "A & B"
    assert data["person_observations"][0]["external_id"] == "001"
    assert all(
        r["origin_raw"] is None and r["event_kind"] == "loyalty_activity"
        for r in data["travel_observations"]
    )
    assert (
        "/cards[1]/card[1]/activities[1]/activity[1]"
        in data["travel_observations"][0]["locator"]
    )
    assert (
        data["travel_observations"][0]["loyalty_identifier_id"]
        == data["identifiers"][0]["identifier_id"]
    )
    assert data["travel_observations"][1]["loyalty_identifier_id"] is None


def test_xml_broken_syntax_and_external_entities(tmp_path):
    path = tmp_path / "bad.xml"
    path.write_text(
        "<PointzAggregatorUsers><user><name>X</name></user>", encoding="utf-8"
    )
    with pytest.raises(Exception):
        list(pointz_xml.read(path))
    secret = tmp_path / "secret"
    secret.write_text("DO_NOT_READ", encoding="utf-8")
    path.write_text(
        f'<!DOCTYPE x [<!ENTITY e SYSTEM "{secret.as_uri()}">]><PointzAggregatorUsers><user><name>&e;</name></user></PointzAggregatorUsers>',
        encoding="utf-8",
    )
    # Entity nodes must be rejected or preserved without expansion, never loaded.
    try:
        data = collect(pointz_xml.read, path)
        assert "DO_NOT_READ" not in json.dumps(data)
    except (ValueError, TypeError):
        pass


def test_sirena_confirmed_character_and_byte_ranges(tmp_path):
    import yaml

    names = [
        ("PaxName", 20),
        ("PaxBirthDate", 10),
        ("DepartDate", 10),
        ("DepartTime", 5),
        ("Flight", 6),
        ("From", 3),
        ("Dest", 3),
        ("e-Ticket", 6),
        ("TravelDoc", 6),
    ]
    fields, start = [], 0
    for name, width in names:
        fields.append(dict(name=name, start=start, end=start + width))
        start += width
    cfg = tmp_path / "columns.yaml"
    cfg.write_text(
        yaml.safe_dump(
            dict(
                confirmed=True,
                encoding="utf-8",
                unit="characters",
                indexing="zero-based-half-open",
                header_lines=1,
                columns=fields,
            )
        ),
        encoding="utf-8",
    )
    values = [
        "Иван Иванов",
        "2000-01-01",
        "2017-01-01",
        "23:30",
        "AB001",
        "",
        "BBB",
        "000123",
        "001234",
    ]
    line = "".join(value.ljust(width) for value, (_, width) in zip(values, names))
    path = tmp_path / "sirena.tab"
    path.write_text("header\n" + line + "\nshort\n" + line + "tail\n", encoding="utf-8")
    data = collect(sirena_fixed.read, path, columns=cfg)
    assert data["person_observations"][0]["name_raw"] == "Иван Иванов"
    assert data["travel_observations"][0]["origin_raw"] is None
    assert data["travel_observations"][0]["ticket_number"] == "000123"
    assert {r["reason"] for r in data["errors"]} == {"short_row", "unexpected_tail"}
    cfg.write_text(
        yaml.safe_dump(
            dict(
                confirmed=True,
                encoding="utf-8",
                unit="bytes",
                indexing="zero-based-half-open",
                header_lines=0,
                columns=[
                    dict(name="PaxName", start=0, end=22),
                    dict(name="e-Ticket", start=22, end=28),
                ],
            )
        ),
        encoding="utf-8",
    )
    path.write_bytes("Иван Иванов".encode().ljust(22) + b"000123\n")
    assert (
        collect(sirena_fixed.read, path, columns=cfg)["identifiers"][0]["value"]
        == "000123"
    )


def test_unconfirmed_sirena_fails(tmp_path):
    path = tmp_path / "x"
    path.write_text("x", encoding="utf-8")
    cfg = tmp_path / "c"
    cfg.write_text("confirmed: false", encoding="utf-8")
    with pytest.raises(ValueError, match="not confirmed"):
        list(sirena_fixed.read(path, columns=cfg))


def test_yaml_zeroes_multiple_cards_neutral_status(tmp_path):
    path = tmp_path / "skyteam.yaml"
    path.write_text(
        "2017-01-01:\n  AB001:\n    FROM: AAA\n    TO: BBB\n    STATUS: UNKNOWN\n    FF:\n      AA: [0012, 0023]\n      BB: 0012\n  AB002:\n    FF: []\n---\n2017-01-02:\n  AB001:\n    FF: [{Programme: AA, Number: 0001}]\n",
        encoding="utf-8",
    )
    data = collect(skyteam_yaml.read, path)
    assert len(data["person_observations"]) == len(data["travel_observations"]) == 4
    assert data["identifiers"][0]["value"] == "0012"
    assert data["identifiers"][2]["programme"] == "BB"
    assert data["travel_observations"][0]["status_raw"] == "UNKNOWN"
    assert data["travel_observations"][0]["event_kind"] == "exchange_record"
    assert "document:1/2017-01-01/AB001/FF/AA/0" in data["identifiers"][0]["locator"]


@pytest.mark.parametrize(
    "source",
    [
        "2017-01-01: {}\n2017-01-01: {}",
        "2017-01-01:\n  AB1: {}\n  AB1: {}",
        "2017-01-01:\n  AB1:\n    FF: {AA: 001, AA: 002}",
        "2017-01-01:\n  AB1: &a {FF: []}\n  AB2: *a",
    ],
)
def test_yaml_duplicates_and_aliases_fail(tmp_path, source):
    path = tmp_path / "skyteam.yaml"
    path.write_text(source, encoding="utf-8")
    with pytest.raises(ValueError):
        ingest(skyteam_yaml.read, path, "skyteam_yaml", tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_xlsx_zip_all_sheets_and_content_date(tmp_path):
    book = tmp_path / "2018-01-01.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Passengers"
    ws.append(
        ["PassengerName", "FlightDate", "FlightNumber", "TicketNumber", "Destination"]
    )
    ws.append(["X", "2017-01-01", "AB001", "000123", "BBB"])
    ws.append(
        ["PassengerName", "FlightDate", "FlightNumber", "TicketNumber", "Destination"]
    )
    ws.append(["Y", "=TODAY()", "AB002", 123, "CCC"])
    wb.create_sheet("Empty")
    wb.create_sheet("Lookup").append(["Airport", "Country"])
    wb.save(book)
    wb.close()
    path = tmp_path / "boarding.zip"
    with ZipFile(path, "w") as z:
        z.write(book, book.name)
        z.writestr("notes.txt", "n")
    cov = tmp_path / "coverage.csv"
    data = collect(boardingpass_xlsx.read, path, coverage_path=cov)
    assert len(data["travel_observations"]) == 2
    assert data["travel_observations"][0]["departure_local"] == "2017-01-01"
    assert data["identifiers"][0]["value"] == "000123"
    assert data["travel_observations"][1]["departure_local"] is None
    reasons = {r["reason"] for r in data["errors"]}
    assert "formula_requires_cached_value" in reasons
    assert "numeric_identifier_original_digits_unavailable" in reasons
    coverage = list(csv.DictReader(cov.open(encoding="utf-8")))
    assert len(coverage) == 4
    assert {r["sheet"] for r in coverage} == {"Passengers", "Empty", "Lookup", ""}
    assert coverage[0]["input_rows"] == "2"
    assert (
        "2018-01-01.xlsx/Passengers/row:2" in data["travel_observations"][0]["locator"]
    )


@pytest.mark.parametrize(
    "source",
    [
        '{"Forum Profiles":{"a":{"Real Name":"X","Real Name":"Y"}}}',
        '{"Forum Profiles":{"a":{"Real Name":"X"},"a":{"Real Name":"Y"}}}',
    ],
)
def test_json_duplicate_fields_do_not_silently_overwrite(tmp_path, source):
    path = tmp_path / "forum.json"
    path.write_text(source, encoding="utf-8")
    with pytest.raises(ValueError, match="Duplicate JSON"):
        list(forum_json.read(path))


def test_xlsx_corrupt_member_does_not_skip_following_books(tmp_path):
    book = tmp_path / "good.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["PassengerName", "FlightDate", "TicketNumber"])
    ws.append(["X", "2017-01-01", "0001"])
    wb.save(book)
    wb.close()
    archive = tmp_path / "books.zip"
    with ZipFile(archive, "w") as z:
        z.writestr("bad.xlsx", b"broken")
        z.write(book, "good.xlsx")
    coverage = tmp_path / "coverage.csv"
    data = collect(boardingpass_xlsx.read, archive, coverage_path=coverage)
    assert len(data["person_observations"]) == 1
    assert "unreadable_workbook" in {r["reason"] for r in data["errors"]}
    with coverage.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert [row["workbook"] for row in rows] == ["bad.xlsx", "good.xlsx"]
    assert rows[0]["status"] == "excluded" and rows[1]["status"] == "ok"


def test_xlsx_uses_cached_formula_when_present(tmp_path):
    from zipfile import ZIP_DEFLATED

    book = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["PassengerName", "FlightDate", "TicketNumber"])
    ws.append(["X", '="2017-01-01"', "0001"])
    wb.save(book)
    wb.close()
    rewritten = tmp_path / "cached.xlsx"
    with ZipFile(book) as src, ZipFile(rewritten, "w", ZIP_DEFLATED) as dst:
        for info in src.infolist():
            value = src.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                value = value.replace(
                    b'<c r="B2"><f>"2017-01-01"</f><v></v></c>',
                    b'<c r="B2" t="str"><f>"2017-01-01"</f><v>2017-01-01</v></c>',
                )
                assert b't="str"' in value
            dst.writestr(info, value)
    archive = tmp_path / "books.zip"
    with ZipFile(archive, "w") as z:
        z.write(rewritten, "cached.xlsx")
    data = collect(boardingpass_xlsx.read, archive)
    assert data["travel_observations"][0]["departure_local"] == "2017-01-01"
    assert not data.get("errors")


def test_xlsx_blank_first_cell_preserves_row_and_invalid_date_coverage(tmp_path):
    book = tmp_path / "blank.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["PassengerName", "FlightDate", "TicketNumber"])
    ws.append([None, "2017-02-30", "000123"])
    wb.save(book)
    wb.close()
    archive = tmp_path / "books.zip"
    with ZipFile(archive, "w") as z:
        z.write(book, "blank.xlsx")
    coverage = tmp_path / "coverage.csv"
    data = collect(boardingpass_xlsx.read, archive, coverage_path=coverage)
    assert data["person_observations"][0]["name_raw"] is None
    assert data["travel_observations"][0]["ticket_number"] == "000123"
    assert data["travel_observations"][0]["locator"].endswith("row:2/travel")
    with coverage.open(encoding="utf-8") as f:
        row = next(csv.DictReader(f))
    assert row["accepted_rows"] == "1" and row["errors"] == "1"
