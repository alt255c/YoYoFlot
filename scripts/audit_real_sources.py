"""Independent selected-record comparison against source values, with locators."""

import csv
import json
import re
import shutil
import tempfile
import hashlib
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile
import duckdb
import ijson
import yaml
from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
import argparse

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--data-dir", type=Path, default=ROOT / "data/raw")
parser.add_argument("--interim", type=Path, default=ROOT / "data/interim/real")
args = parser.parse_args()
INPUT = args.data_dir
INTERIM = args.interim
OUT = ROOT / "docs/validation"
census = json.loads((OUT / "input_census.json").read_text(encoding="utf-8"))
checks = []


def digest(value):
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, default=str).encode()
    ).hexdigest()


def check(source, locator, field, expected, actual):
    checks.append(
        {
            "source": source,
            "locator": locator,
            "field": field,
            "match": expected == actual,
            "expected_sha256": digest(expected),
            "actual_sha256": digest(actual),
        }
    )
    assert expected == actual, (source, locator, field, expected, actual)


def connection(source):
    db = duckdb.connect()
    db.execute("SET memory_limit='512MB'")
    db.execute("SET threads=2")
    for name in ["person_observations", "identifiers", "travel_observations"]:
        db.from_parquet(str(INTERIM / source / (name + ".parquet"))).create_view(name)
    return db


def one(db, table, locator):
    rows = db.execute(f"SELECT * FROM {table} WHERE locator=?", [locator]).fetchall()
    names = [d[0] for d in db.description]
    assert len(rows) == 1, (table, locator, len(rows))
    return dict(zip(names, rows[0]))


# CSV: 20 records spread across the full input, compare dates using confirmed MDY.
targets = {1 + (census["boarding_csv"]["valid_width"] - 1) * i // 19 for i in range(20)}
with (
    connection("boarding_csv") as db,
    (INPUT / "BoardingData.csv").open(encoding="utf-8-sig", newline="") as f,
):
    r = csv.reader(f, delimiter=";")
    header = next(r)
    index = 0
    while True:
        first = r.line_num + 1
        try:
            row = next(r)
        except StopIteration:
            break
        if len(row) != 14:
            continue
        index += 1
        if index not in targets:
            continue
        obj = dict(zip(header, row))
        loc = f"line:{first}-{r.line_num}"
        p = one(db, "person_observations", loc)
        t = one(db, "travel_observations", loc + "/travel")
        check(
            "boarding_csv",
            loc,
            "birth_date",
            datetime.strptime(obj["PassengerBirthDate"], "%m/%d/%Y").date().isoformat(),
            p["birth_date"],
        )
        check("boarding_csv", loc, "origin", None, t["origin_raw"])
        check(
            "boarding_csv",
            loc,
            "flight_number",
            obj["FlightNumber"],
            t["flight_number"],
        )
        check(
            "boarding_csv", loc, "destination", obj["Destination"], t["destination_raw"]
        )
        check(
            "boarding_csv",
            loc,
            "ticket",
            None
            if obj["TicketNumber"] in ("Not presented", "None", "")
            else obj["TicketNumber"],
            t["ticket_number"],
        )

# Sirena: all 20 columns from independent header matching, 20 dispersed rows.
with (INPUT / "Sirena-export-fixed.tab").open(encoding="utf-8") as f:
    header = f.readline().rstrip("\n")
    names = [
        "PaxName",
        "PaxBirthDate",
        "DepartDate",
        "DepartTime",
        "ArrivalDate",
        "ArrivalTime",
        "Flight",
        "CodeSh",
        "From",
        "Dest",
        "Code",
        "e-Ticket",
        "TravelDoc",
        "Seat",
        "Meal",
        "TrvCls",
        "Fare",
        "Baggage",
        "PaxAdditionalInfo",
        "AgentInfo",
    ]
    starts = []
    pos = 0
    for name in names:
        pos = header.index(name, pos)
        starts.append(pos)
        pos += len(name)
    ends = starts[1:] + [len(header)]
    target = {2 + (census["sirena_fixed"]["rows"] - 1) * i // 19 for i in range(20)}
    audit = []
    with connection("sirena_fixed") as db:
        for line_number, line in enumerate(f, 2):
            if line_number not in target:
                continue
            obj = {
                name: line[s:e].strip() or None
                for name, s, e in zip(names, starts, ends)
            }
            loc = f"line:{line_number}"
            p = one(db, "person_observations", loc)
            t = one(db, "travel_observations", loc + "/travel")
            raw = json.loads(p["raw_json"])
            for name in names:
                check("sirena_fixed", loc, "raw:" + name, obj[name], raw[name])
            for source_field, target_field in [
                ("Flight", "flight_number"),
                ("From", "origin_raw"),
                ("Dest", "destination_raw"),
                ("e-Ticket", "ticket_number"),
                ("Code", "booking_code"),
                ("CodeSh", "codeshare"),
            ]:
                check(
                    "sirena_fixed",
                    loc,
                    target_field,
                    obj[source_field],
                    t[target_field],
                )
            check("sirena_fixed", loc, "name", obj["PaxName"], p["name_raw"])
            check(
                "sirena_fixed",
                loc,
                "birth_date",
                None if obj["PaxBirthDate"] == "N/A" else obj["PaxBirthDate"],
                p["birth_date"],
            )
            for prefix in ["departure", "arrival"]:
                source_prefix = "Depart" if prefix == "departure" else "Arrival"
                check(
                    "sirena_fixed",
                    loc,
                    prefix,
                    obj[source_prefix + "Date"] + "T" + obj[source_prefix + "Time"],
                    t[prefix + "_local"],
                )
            audit.append(
                {
                    "locator": loc,
                    **obj,
                    "match": "yes",
                    "review_method": "source/output field comparison",
                }
            )
    with (OUT / "sirena_audit_20.csv").open("w", encoding="utf-8", newline="") as out:
        w = csv.DictWriter(out, fieldnames=list(audit[0]))
        w.writeheader()
        w.writerows(audit)

# XML: 5 users across file, all names/cards and selected nested activity links.
target = {1 + (census["pointz_xml"]["user"] - 1) * i // 4 for i in range(5)}
index = 0
with connection("pointz_xml") as db:
    for _, user in etree.iterparse(
        str(INPUT / "PointzAggregator-AirlinesData.xml"),
        events=("end",),
        tag="user",
        resolve_entities=False,
        no_network=True,
    ):
        index += 1
        if index in target:
            loc = f"/PointzAggregatorUsers/user[{index}]"
            p = one(db, "person_observations", loc)
            n = user.find("name")
            check(
                "pointz_xml",
                loc,
                "name",
                " ".join(n.get(k) for k in ["first", "middle", "last"] if n.get(k)),
                p["name_raw"],
            )
            check("pointz_xml", loc, "uid", user.get("uid"), p["external_id"])
            cards = user.findall("cards/card")
            for i, card in enumerate(cards, 1):
                cloc = loc + f"/cards/card[{i}]"
                row = one(db, "identifiers", cloc)
                programme, number = card.get("number").split()
                check("pointz_xml", cloc, "programme", programme, row["programme"])
                check("pointz_xml", cloc, "number", number, row["value"])
                activities = card.findall("activities/activity")
                for ai in sorted({1, len(activities)}):
                    if not activities:
                        continue
                    activity = activities[ai - 1]
                    aloc = loc + f"/cards[1]/card[{i}]/activities[1]/activity[{ai}]"
                    t = one(db, "travel_observations", aloc)
                    check(
                        "pointz_xml",
                        aloc,
                        "card_link",
                        row["identifier_id"],
                        t["loyalty_identifier_id"],
                    )
                    check(
                        "pointz_xml",
                        aloc,
                        "date",
                        activity.findtext("Date"),
                        t["departure_local"],
                    )
                    check(
                        "pointz_xml",
                        aloc,
                        "from",
                        activity.findtext("Departure"),
                        t["origin_raw"],
                    )
                    check(
                        "pointz_xml",
                        aloc,
                        "to",
                        activity.findtext("Arrival"),
                        t["destination_raw"],
                    )
        user.clear(keep_tail=True)
        while user.getprevious() is not None:
            del user.getparent()[0]

# JSON: independent ijson traversal, one profile at a time; no documents in this input.
target = {(census["forum_json"]["profiles"] - 1) * i // 4 for i in range(5)}
with (
    connection("forum_json") as db,
    (INPUT / "FrequentFlyerForum-Profiles.json").open("rb") as f,
):
    for index, obj in enumerate(ijson.items(f, "Forum Profiles.item")):
        if index not in target:
            continue
        loc = f"/Forum Profiles/{index}"
        p = one(db, "person_observations", loc)
        name = (
            " ".join(
                obj["Real Name"][k]
                for k in ["First Name", "Middle Name", "Last Name"]
                if obj["Real Name"].get(k)
            )
            or None
        )
        check("forum_json", loc, "name", name, p["name_raw"])
        travels = obj["Registered Flights"]
        count = db.execute(
            "SELECT count(*) FROM travel_observations WHERE observation_id=?",
            [p["observation_id"]],
        ).fetchone()[0]
        check("forum_json", loc, "travel_count", len(travels), count)
        for i, card in enumerate(obj["Loyality Programm"]):
            cloc = loc + f"/Loyality Programm/{i}"
            row = one(db, "identifiers", cloc)
            check("forum_json", cloc, "programme", card["programm"], row["programme"])
            check("forum_json", cloc, "number", card["Number"].strip(), row["value"])
        for i in sorted({0, len(travels) - 1}):
            if not travels:
                continue
            t = one(db, "travel_observations", loc + f"/Registered Flights/{i}")
            flight = travels[i]
            check(
                "forum_json",
                t["locator"],
                "codeshare",
                str(flight["Codeshare"]),
                t["codeshare"],
            )
            check(
                "forum_json", t["locator"], "date", flight["Date"], t["departure_local"]
            )
            check(
                "forum_json",
                t["locator"],
                "origin",
                flight["Departure"],
                json.loads(t["origin_raw"]),
            )
            check(
                "forum_json",
                t["locator"],
                "destination",
                flight["Arrival"],
                json.loads(t["destination_raw"]),
            )

# YAML: extract 5 literal date/flight blocks, parse only these small blocks independently.
selected = []
date_index = -1
day = None
flight = None
lines = []
targets = {0, 91, 182, 274, 365}
take_day = False
with (INPUT / "SkyTeam-Exchange.yaml").open(encoding="utf-8-sig") as f:
    for line in f:
        date_match = re.match(r"^'([^']+)':", line)
        flight_match = re.match(r"^  ([^ ]+):\s*$", line)
        if date_match or flight_match:
            if lines:
                selected.append((day, flight, lines))
                lines = []
                take_day = False
        if date_match:
            date_index += 1
            day = date_match[1]
            take_day = date_index in targets
        elif flight_match:
            flight = flight_match[1]
            if take_day:
                lines = [line]
        elif lines:
            lines.append(line)
    if lines:
        selected.append((day, flight, lines))
with connection("skyteam_yaml") as db:
    for day, flight, lines in selected:
        obj = yaml.load("".join(lines), Loader=yaml.BaseLoader)[flight]
        for key, payload in obj["FF"].items():
            programme, number = key.split()
            loc = f"document:1/{day}/{flight}/FF/{key}"
            i = one(db, "identifiers", loc + "/card")
            t = one(db, "travel_observations", loc + "/travel")
            for field, expected in [("programme", programme), ("value", number)]:
                check("skyteam_yaml", loc, field, expected, i[field])
            for field, expected in [
                ("flight_number", flight),
                ("departure_local", day),
                ("origin_raw", obj["FROM"]),
                ("destination_raw", obj["TO"]),
                ("status_raw", obj["STATUS"]),
                ("fare_raw", payload["FARE"]),
            ]:
                check("skyteam_yaml", loc, field, expected, t[field])

# XLSX: independent openpyxl reading of first/middle/end/largest books, 2 sheets each.
with (
    connection("boardingpass_xlsx") as db,
    ZipFile(INPUT / "YourBoardingPassDotAero.zip") as z,
):
    books = z.infolist()
    chosen = {
        0,
        len(books) // 2,
        len(books) - 1,
        max(range(len(books)), key=lambda i: books[i].file_size),
    }
    for bi in sorted(chosen):
        member = books[bi]
        with z.open(member) as src, tempfile.TemporaryFile() as tmp:
            shutil.copyfileobj(src, tmp, 1024 * 1024)
            tmp.seek(0)
            wb = load_workbook(tmp, read_only=True, data_only=True)
            try:
                for si in sorted({0, len(wb.sheetnames) - 1}):
                    sheet = wb.worksheets[si]
                    cells = {
                        c.coordinate: c.value
                        for row in sheet
                        for c in row
                        if c.value is not None
                    }
                    loc = f"YourBoardingPassDotAero.zip/{member.filename}/{sheet.title}/row:3"
                    p = one(db, "person_observations", loc)
                    t = one(db, "travel_observations", loc + "/travel")
                    check("boardingpass_xlsx", loc, "name", cells["B3"], p["name_raw"])
                    for field, coord in [
                        ("flight_number", "A5"),
                        ("origin_raw", "D7"),
                        ("destination_raw", "H7"),
                        ("ticket_number", "E13"),
                        ("booking_code", "B13"),
                    ]:
                        check(
                            "boardingpass_xlsx",
                            loc,
                            field,
                            None if cells.get(coord) in (None, "N/A") else cells[coord],
                            t[field],
                        )
                    check(
                        "boardingpass_xlsx",
                        loc,
                        "departure",
                        cells["A9"] + "T" + cells["C9"],
                        t["departure_local"],
                    )
            finally:
                wb.close()
(OUT / "source_audit_checks.json").write_text(
    json.dumps(
        {
            "method": "independent source/output comparison; assistant review, not a human sign-off",
            "checks": checks,
            "count": len(checks),
            "failed": sum(not x["match"] for x in checks),
        },
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)
print("Source/output checks:", len(checks), "passed")
